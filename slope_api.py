"""PLS GIS v1.2 - slope clipping backend. Run: uvicorn app:app --host 0.0.0.0 --port 8000"""
import os, io, json, zipfile, tempfile, shutil, uuid
from pathlib import Path
from functools import lru_cache
import geopandas as gpd
import pandas as pd
from shapely.geometry import shape, mapping
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

ROOT=Path(__file__).resolve().parent
DATA=Path(os.environ.get('PLS_DATA_DIR',str(ROOT/'data')))
SLOPE=Path(os.environ.get('PLS_SLOPE_SHP',str(DATA/'slope'/'Slope_utm.shp')))
BOUNDARIES=Path(os.environ.get('PLS_BOUNDARY_FILE',str(DATA/'admin_boundaries.geojson')))
OUT=Path(os.environ.get('PLS_OUTPUT_DIR',str(ROOT/'outputs')));OUT.mkdir(parents=True,exist_ok=True)
app=FastAPI(title='Philippine Land Services GIS API',version='1.2')
origins=[s.strip() for s in os.environ.get('PLS_CORS_ORIGINS','http://localhost:5500,http://127.0.0.1:5500,http://localhost:8000').split(',') if s.strip()]
app.add_middleware(CORSMiddleware,allow_origins=origins,allow_methods=['GET','POST'],allow_headers=['*'])
FIELDS=['SLOPE','FIRST_DESC','FIRST_CODE']
COLORS={'M':'#2d9d5b','N':'#a5ce55','O':'#f5e46b','P':'#f5a24b','Q':'#dc6346','R':'#a43b64','RIVER':'#54a8d9'}

@app.get('/api/health')
def health():return {'status':'ok','slope_ready':SLOPE.exists(),'boundaries_ready':BOUNDARIES.exists(),'version':'1.2'}

@lru_cache(maxsize=1)
def slope_gdf():
 if not SLOPE.exists():raise FileNotFoundError(f'Slope database missing: {SLOPE}')
 g=gpd.read_file(SLOPE)
 if g.crs is None:raise ValueError('Slope database has no CRS')
 return g

@lru_cache(maxsize=1)
def admin_gdf():
 if not BOUNDARIES.exists():return None
 g=gpd.read_file(BOUNDARIES)
 if g.crs is None:raise ValueError('Boundary catalog missing CRS')
 return g.to_crs(4326)

def admin_field(g,role):
 for c in g.columns:
  if c.lower() in {'region':['region','reg_name','adm1_en','adm1_name','regname'],'province':['province','prov_name','adm2_en','adm2_name','provname'],'municipality':['municipality','mun_name','citymun','adm3_en','adm3_name','city_mun','munname']}[role]:return c
 return None

@app.get('/api/regions')
def regions():
 g=admin_gdf()
 if g is None:return {'configured':False,'regions':[]}
 f=admin_field(g,'region');return {'configured':bool(f),'regions':sorted(g[f].dropna().astype(str).unique().tolist()) if f else []}

@app.get('/api/provinces')
def provinces(region:str):
 g=admin_gdf();rf=admin_field(g,'region') if g is not None else None;pf=admin_field(g,'province') if g is not None else None
 if not rf or not pf:return {'provinces':[]}
 return {'provinces':sorted(g.loc[g[rf].astype(str)==region,pf].dropna().astype(str).unique().tolist())}

@app.get('/api/municipalities')
def municipalities(region:str,province:str):
 g=admin_gdf();rf=admin_field(g,'region') if g is not None else None;pf=admin_field(g,'province') if g is not None else None;mf=admin_field(g,'municipality') if g is not None else None
 if not all([rf,pf,mf]):return {'municipalities':[]}
 s=g.loc[(g[rf].astype(str)==region)&(g[pf].astype(str)==province),mf]
 return {'municipalities':sorted(s.dropna().astype(str).unique().tolist())}

@app.get('/api/classes')
def classes():
 g=slope_gdf();v=g.groupby(FIELDS,dropna=False).size().reset_index(name='features')
 return {'crs':str(g.crs),'feature_count':len(g),'classes':[{'slope':None if pd.isna(r.SLOPE) else str(r.SLOPE),'description':str(r.FIRST_DESC),'code':str(r.FIRST_CODE),'features':int(r.features),'color':COLORS.get(str(r.FIRST_CODE),'#aaa')} for r in v.itertuples(index=False)]}

def safe_zip_extract(blob,dest):
 with zipfile.ZipFile(io.BytesIO(blob)) as z:
  names=z.namelist()
  if len(names)>100:raise ValueError('Too many files in ZIP')
  for info in z.infolist():
   p=Path(info.filename)
   if p.is_absolute() or '..' in p.parts:raise ValueError('Unsafe ZIP path')
   if info.file_size>250*1024*1024:raise ValueError('ZIP member too large')
   if p.suffix.lower() not in {'.shp','.shx','.dbf','.prj','.cpg','.sbn','.sbx','.xml',''}:raise ValueError('Unsupported ZIP file')
  z.extractall(dest)
 shp=list(dest.rglob('*.shp'))
 if len(shp)!=1:raise ValueError('ZIP must contain exactly one shapefile')
 return shp[0]

def boundary_from_upload(upload,blob,work):
 suffix=Path(upload.filename or '').suffix.lower()
 if suffix=='.zip':g=gpd.read_file(safe_zip_extract(blob,work))
 elif suffix in ('.json','.geojson'):
  p=work/'boundary.geojson';p.write_bytes(blob);g=gpd.read_file(p)
 else:raise ValueError('Upload ZIP shapefile or GeoJSON')
 if g.crs is None:raise ValueError('Uploaded boundary lacks CRS; include .prj')
 return g

def selected_boundary(region,province,municipality):
 g=admin_gdf()
 if g is None:raise ValueError('Administrative boundary catalog not installed. Upload a parcel boundary instead.')
 rf,pf,mf=[admin_field(g,k) for k in ('region','province','municipality')]
 if not all([rf,pf,mf]):raise ValueError('Boundary catalog must have region, province, municipality fields')
 if not all([region,province,municipality]):raise ValueError('Choose region, province and municipality')
 g=g[(g[rf].astype(str)==region)&(g[pf].astype(str)==province)&(g[mf].astype(str)==municipality)]
 if g.empty:raise ValueError('Municipality not found')
 return g

def clip_slope(boundary):
 if boundary.empty:raise ValueError('Empty boundary')
 boundary=boundary[boundary.geometry.notna() & ~boundary.geometry.is_empty].copy()
 if not boundary.geom_type.isin(['Polygon','MultiPolygon']).all():raise ValueError('Boundary must contain only polygons')
 boundary=boundary.to_crs(slope_gdf().crs)
 boundary['geometry']=boundary.geometry.make_valid()
 geom=boundary.geometry.union_all()
 source=slope_gdf()
 candidates=source.iloc[list(source.sindex.query(geom,predicate='intersects'))].copy()
 if candidates.empty:raise ValueError('No slope features intersect the boundary')
 candidates['geometry']=candidates.geometry.make_valid()
 clipped=gpd.clip(candidates,gpd.GeoDataFrame(geometry=[geom],crs=source.crs),keep_geom_type=True)
 clipped=clipped[clipped.geometry.notna() & ~clipped.geometry.is_empty].copy()
 clipped['AREA_HA']=(clipped.geometry.area/10000).round(3)
 clipped=clipped[clipped.geometry.area>0].copy()
 if clipped.empty:raise ValueError('No polygon area in selected boundary')
 return clipped

def make_report(g,out_dir):
 summary=g.groupby(FIELDS,dropna=False,sort=False)['AREA_HA'].sum().reset_index()
 summary['AREA_HA']=summary['AREA_HA'].round(3)
 total=float(summary['AREA_HA'].sum());summary['PERCENT']=((summary['AREA_HA']/total)*100).round(2) if total else 0
 wb=Workbook();ws=wb.active;ws.title='Slope Summary';ws.append(['SLOPE','DESCRIPTION','CODE','AREA_HA','PERCENT'])
 for r in summary.itertuples(index=False):ws.append([str(r.SLOPE) if pd.notna(r.SLOPE) else '',str(r.FIRST_DESC),str(r.FIRST_CODE),float(r.AREA_HA),float(r.PERCENT)])
 ws.append(['TOTAL','','',total,100 if total else 0]);ws.freeze_panes='A2'
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='102D4B')
 for col,w in {'A':18,'B':38,'C':18,'D':16,'E':16}.items():ws.column_dimensions[col].width=w
 wb.save(out_dir/'slope_summary.xlsx')
 rows=[['Slope','Description','Code','Area (ha)','%']]+[[str(r.SLOPE) if pd.notna(r.SLOPE) else '—',str(r.FIRST_DESC),str(r.FIRST_CODE),f'{r.AREA_HA:,.3f}',f'{r.PERCENT:.2f}'] for r in summary.itertuples(index=False)]
 doc=SimpleDocTemplate(str(out_dir/'slope_report.pdf'),pagesize=(842,595),leftMargin=35,rightMargin=35)
 styles=getSampleStyleSheet();t=Table(rows,repeatRows=1,colWidths=[100,300,90,100,65]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#102d4b')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.4,colors.lightgrey),('VALIGN',(0,0),(-1,-1),'TOP'),('FONTSIZE',(0,0),(-1,-1),9)]))
 doc.build([Paragraph('Philippine Land Services — Slope Area Report',styles['Title']),Spacer(1,14),Paragraph(f'Clipped area: {total:,.3f} hectares. Source: Slope_utm (EPSG:32651).',styles['Normal']),Spacer(1,16),t])
 return summary

@app.post('/api/slope/clip')
async def slope_clip(parcel:UploadFile|None=File(None),region:str=Form(''),province:str=Form(''),municipality:str=Form('')):
 try:
  job=uuid.uuid4().hex;dest=OUT/job;dest.mkdir(parents=True,exist_ok=True)
  with tempfile.TemporaryDirectory() as tmp:
   if parcel:
    blob=await parcel.read()
    if len(blob)>100*1024*1024:raise ValueError('Parcel upload exceeds 100 MB')
    boundary=boundary_from_upload(parcel,blob,Path(tmp))
   else:boundary=selected_boundary(region,province,municipality)
   clipped=clip_slope(boundary)
   clipped.to_file(dest/'slope_clipped.shp',driver='ESRI Shapefile',encoding='UTF-8')
   clipped.to_crs(4326).to_file(dest/'slope_clipped.geojson',driver='GeoJSON')
   summary=make_report(clipped,dest)
   with zipfile.ZipFile(dest/'slope_results.zip','w',zipfile.ZIP_DEFLATED) as z:
    for p in dest.iterdir():
     if p.is_file() and p.name!='slope_results.zip':z.write(p,p.name)
  return {'job_id':job,'feature_count':len(clipped),'area_ha':round(float(clipped.geometry.area.sum()/10000),3),'summary':[{'slope':None if pd.isna(r.SLOPE) else str(r.SLOPE),'description':str(r.FIRST_DESC),'code':str(r.FIRST_CODE),'area_ha':float(r.AREA_HA),'percent':float(r.PERCENT)} for r in summary.itertuples(index=False)],'download_url':f'/api/download/{job}/slope_results.zip','preview_url':f'/api/download/{job}/slope_clipped.geojson'}
 except (ValueError,FileNotFoundError,zipfile.BadZipFile) as e:raise HTTPException(400,str(e))

@app.get('/api/download/{job}/{filename}')
def download(job:str,filename:str):
 allowed={'slope_results.zip','slope_clipped.geojson','slope_summary.xlsx','slope_report.pdf'}
 if filename not in allowed or not job.isalnum():raise HTTPException(404,'Not found')
 p=OUT/job/filename
 if not p.is_file():raise HTTPException(404,'Not found')
 return FileResponse(p,filename=filename)
